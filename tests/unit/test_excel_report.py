"""Unit tests for Layer 13 — stof.reporting.excel_report."""
from datetime import datetime, timezone

from openpyxl import load_workbook

from stof.crawler.endpoint_store import Endpoint
from stof.findings.models import Finding
from stof.reporting.excel_report import write


def _finding(**overrides) -> Finding:
    endpoint = Endpoint(url="https://x/admin/admin.jsp", method="GET", endpoint_type="page", parameters=[])
    defaults = dict(
        module_id="idor_tests", vuln_type="Vertical Privilege Escalation", severity="High", cvss_score=8.8,
        endpoint=endpoint, user_role="normal", request_raw="GET x", response_raw="HTTP 200",
        description="desc", recommendation="rec", discovered_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
    )
    defaults.update(overrides)
    return Finding(**defaults)


def test_write_produces_a_readable_workbook_with_header_and_rows(tmp_path):
    findings = [_finding(finding_id="f1"), _finding(finding_id="f2", severity="High")]
    path = write(findings, tmp_path / "report.xlsx")

    wb = load_workbook(path)
    ws = wb.active

    assert ws.cell(row=1, column=1).value == "Finding ID"
    assert ws.cell(row=2, column=1).value == "f1"
    assert ws.cell(row=3, column=1).value == "f2"
    assert ws.cell(row=2, column=3).value == "High"


def test_write_with_no_findings_still_produces_header_only(tmp_path):
    path = write([], tmp_path / "empty.xlsx")

    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(row=1, column=1).value == "Finding ID"
    assert ws.cell(row=2, column=1).value is None


def test_write_creates_parent_directories(tmp_path):
    path = write([], tmp_path / "nested" / "dir" / "report.xlsx")
    assert path.is_file()


def test_write_includes_evidence_refs_joined(tmp_path):
    finding = _finding(evidence_refs=["a.png", "b.json"])
    path = write([finding], tmp_path / "report.xlsx")

    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(row=2, column=13).value == "a.png; b.json"


def test_write_includes_scanner_source_column(tmp_path):
    findings = [_finding(finding_id="f1", scanner_source="stof"), _finding(finding_id="f2", scanner_source="burp")]
    path = write(findings, tmp_path / "report.xlsx")

    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(row=1, column=9).value == "Source"
    assert ws.cell(row=2, column=9).value == "STOF"
    assert ws.cell(row=3, column=9).value == "Burp Suite Pro"


def test_severity_score_column_is_honestly_labeled_not_cvss(tmp_path):
    """Labeling the score column "CVSS Score" with no disclaimer would
    let a client mistake it for a real calculated score with nothing
    backing it -- e.g. feeding it into an SLA. The column must be
    honestly named, point at the adjacent real vector column, and the
    disclaimer must be visible on the sheet itself, not just in a
    hover-only comment."""
    finding = _finding(finding_id="f1")
    path = write([finding], tmp_path / "report.xlsx")

    wb = load_workbook(path)
    ws = wb.active
    header_cell = ws.cell(row=1, column=4)
    assert header_cell.value == "Severity Score (Illustrative)"
    assert "CVSS" not in header_cell.value
    assert header_cell.comment is not None
    assert "CVSS v3.1 Vector column" in header_cell.comment.text
    assert ws.cell(row=1, column=5).value == "CVSS v3.1 Vector"

    # An always-visible footer note, not just the hover comment.
    found_note = any(
        row[0].value and "CVSS v3.1 Vector column" in str(row[0].value)
        for row in ws.iter_rows(min_row=2, max_col=1)
    )
    assert found_note, "expected an always-visible disclaimer row below the data"


def test_write_includes_cvss_vector_column(tmp_path):
    finding = _finding(finding_id="f1")
    finding.cvss_vector = "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:N/I:H/A:H"
    path = write([finding], tmp_path / "report.xlsx")

    wb = load_workbook(path)
    ws = wb.active
    assert ws.cell(row=2, column=5).value == "CVSS:3.1/AV:N/AC:L/PR:L/UI:N/S:U/C:N/I:H/A:H"


def test_write_leaves_cvss_vector_blank_when_none(tmp_path):
    finding = _finding(finding_id="f1")
    finding.cvss_vector = None
    path = write([finding], tmp_path / "report.xlsx")

    wb = load_workbook(path)
    ws = wb.active
    assert not ws.cell(row=2, column=5).value  # openpyxl normalizes an empty string cell to None on read-back
