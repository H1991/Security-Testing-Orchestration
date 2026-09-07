"""Unit tests for Layer 13 — stof.reporting.reporter."""
import json
from datetime import datetime, timezone

from openpyxl import load_workbook

from stof.crawler.endpoint_store import Endpoint
from stof.findings.models import Finding
from stof.reporting.reporter import generate_reports


def _finding(**overrides) -> Finding:
    endpoint = Endpoint(url="https://x/bank/showAccount", method="GET", endpoint_type="api", parameters=["listAccounts"])
    defaults = dict(
        module_id="idor_tests", vuln_type="IDOR", severity="High", cvss_score=8.1, endpoint=endpoint,
        user_role="admin", request_raw="GET x", response_raw="HTTP 200", description="d", recommendation="r",
        discovered_at=datetime(2026, 1, 1, tzinfo=timezone.utc),
    )
    defaults.update(overrides)
    return Finding(**defaults)


def test_generate_reports_writes_all_three_formats(tmp_path):
    findings = [_finding()]
    metadata = {"scan_id": "abc123", "target": "https://x", "modules_run": ["idor_tests"]}

    paths = generate_reports(findings, metadata, output_dir=tmp_path)

    assert paths.json.is_file()
    assert paths.html.is_file()
    assert paths.excel.is_file()


def test_generate_reports_file_names_use_scan_id(tmp_path):
    paths = generate_reports([], {"scan_id": "xyz"}, output_dir=tmp_path)

    assert paths.json.name == "scan_xyz.json"
    assert paths.html.name == "scan_xyz.html"
    assert paths.excel.name == "scan_xyz.xlsx"


def test_generate_reports_defaults_scan_id_when_missing(tmp_path):
    paths = generate_reports([], {}, output_dir=tmp_path)

    assert paths.json.name == "scan_scan.json"


def test_generate_reports_content_is_consistent_across_formats(tmp_path):
    findings = [_finding(finding_id="f1"), _finding(finding_id="f2", severity="High")]
    paths = generate_reports(findings, {"scan_id": "abc"}, output_dir=tmp_path)

    json_data = json.loads(paths.json.read_text(encoding="utf-8"))
    assert json_data["summary"]["total_findings"] == 2

    wb = load_workbook(paths.excel)
    ws = wb.active
    assert ws.cell(row=2, column=1).value == "f1"
    assert ws.cell(row=3, column=1).value == "f2"


def test_generate_reports_threads_recon_report_into_json_and_html(tmp_path):
    recon = {"target": "https://x", "pages_analyzed": 4, "tech_stack": [{"url": "https://x/", "tech": ["Java"]}]}

    paths = generate_reports([], {"scan_id": "abc"}, output_dir=tmp_path, recon_report=recon)

    json_data = json.loads(paths.json.read_text(encoding="utf-8"))
    assert json_data["recon"] == recon

    html = paths.html.read_text(encoding="utf-8")
    assert "Reconnaissance" in html
    assert "Java" in html
